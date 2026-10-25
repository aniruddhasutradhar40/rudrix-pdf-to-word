"""Editable Word mode: rebuilds the document with python-docx from PyMuPDF structure.

Text becomes real paragraphs/headings, tables become real Word tables, images become inline
pictures. Page size and margins follow the PDF. This favours editing over pixel fidelity.
"""
from __future__ import annotations

import io
import re
import statistics
from collections import Counter

import pymupdf
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

BENGALI_RE = re.compile(r"[\u0980-\u09FF]")
CTRL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")
BULLET_RE = re.compile(r"^\s*([•●▪◦○■□\-–—*·])\s+")
NUMBER_RE = re.compile(r"^\s*(\(?\d{1,3}[.)]|\(?[a-zA-Z][.)])\s+")
TEXT_FLAGS = pymupdf.TEXT_MEDIABOX_CLIP
BN_FONT = "Nirmala UI"
GENERIC_FONTS = ("glyphless", "tesseract", "unknown", "")


def clean(text: str) -> str:
    return CTRL_RE.sub("", text)


def map_font(name: str) -> str:
    base = name.split("+")[-1]
    base = re.split(r"[-,]", base)[0]
    base = re.sub(r"(PSMT|MT|PS)$", "", base)
    if base.lower() in GENERIC_FONTS:
        return "Calibri"
    return re.sub(r"(?<=[a-z])(?=[A-Z])", " ", base) if " " not in base else base


def set_run_fonts(run, font: str, bengali: bool) -> None:
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    f = BN_FONT if bengali else font
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), f)
    if bengali:
        lang = OxmlElement("w:lang")
        lang.set(qn("w:bidi"), "bn-BD")
        rpr.append(lang)


def shrink_paragraph(p) -> None:
    pf = p.paragraph_format
    pf.space_before = pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(1)
    ppr = p._p.get_or_add_pPr()
    rpr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "2")
    rpr.append(sz)
    ppr.append(rpr)


def body_font_size(page_dicts: list[dict]) -> float:
    sizes: Counter = Counter()
    for d in page_dicts:
        for b in d["blocks"]:
            if b["type"] != 0:
                continue
            for ln in b["lines"]:
                for sp in ln["spans"]:
                    n = len(sp["text"].strip())
                    if n:
                        sizes[round(sp["size"] * 2) / 2] += n
    return sizes.most_common(1)[0][0] if sizes else 11.0


def column_order(items: list[dict], page_w: float) -> list[dict]:
    """Reading order. Detects two-column pages; otherwise top-to-bottom, left-to-right."""
    mid = page_w / 2
    def side(it):
        if it["x1"] <= mid + 4:
            return 0
        if it["x0"] >= mid - 4:
            return 1
        return 2
    texty = [i for i in items if i["kind"] == "text"]
    left = [i for i in texty if side(i) == 0 and i["lines"] >= 2]
    right = [i for i in texty if side(i) == 1 and i["lines"] >= 2]
    ordered = sorted(items, key=lambda i: (round(i["y0"] / 4), i["x0"]))
    if len(left) < 2 or len(right) < 2:
        return ordered
    out, band = [], ([], [])
    def flush():
        out.extend(sorted(band[0], key=lambda i: i["y0"]))
        out.extend(sorted(band[1], key=lambda i: i["y0"]))
        band[0].clear(); band[1].clear()
    for it in sorted(items, key=lambda i: i["y0"]):
        s = side(it)
        if s == 2:
            flush(); out.append(it)
        else:
            band[s].append(it)
    flush()
    return out


def extract_image_bytes(doc: pymupdf.Document, xref: int) -> bytes | None:
    try:
        pix = pymupdf.Pixmap(doc, xref)
        smask = doc.xref_get_key(xref, "SMask")
        if smask[0] == "xref":
            mask = pymupdf.Pixmap(doc, int(smask[1].split()[0]))
            if pix.alpha:
                pix = pymupdf.Pixmap(pix, 0)
            if pix.colorspace and pix.colorspace.n != 3:
                pix = pymupdf.Pixmap(pymupdf.csRGB, pix)
            pix = pymupdf.Pixmap(pix, mask)
        elif pix.colorspace is None or pix.colorspace.n not in (1, 3) or pix.alpha:
            if pix.alpha and pix.colorspace and pix.colorspace.n in (1, 3):
                pass
            else:
                pix = pymupdf.Pixmap(pymupdf.csRGB, pix)
        return pix.tobytes("png")
    except Exception:
        return None


class Builder:
    def __init__(self, pdf_path: str, *, ocr_lang: str | None, ocr_tessdata: str | None, ocr_max_pages: int, scanned_pages: set[int]):
        self.doc_pdf = pymupdf.open(pdf_path)
        self.ocr_lang, self.tessdata, self.ocr_max = ocr_lang, ocr_tessdata, ocr_max_pages
        self.scanned = scanned_pages
        self.out = Document()
        self.warnings: list[str] = []
        self.stats = {"pages": 0, "paragraphs": 0, "headings": 0, "tables": 0, "images": 0, "ocr_pages": 0, "image_only_pages": 0}
        self.ocr_pages: list[int] = []
        st = self.out.styles["Normal"]
        st.font.name, st.font.size = "Calibri", Pt(11)

    # ---- page data -------------------------------------------------------
    def page_dict(self, i: int, page) -> tuple[dict, bool]:
        if i in self.scanned and self.ocr_lang and len(self.ocr_pages) < self.ocr_max:
            try:
                tp = page.get_textpage_ocr(language=self.ocr_lang, dpi=300, full=True, tessdata=self.tessdata)
                self.ocr_pages.append(i + 1)
                return page.get_text("dict", flags=TEXT_FLAGS, textpage=tp), True
            except Exception as exc:  # OCR failed -> image fallback
                self.warnings.append(f"OCR failed on page {i + 1} ({type(exc).__name__}); inserted the page as an image.")
        return page.get_text("dict", flags=TEXT_FLAGS), False

    def build(self, out_path: str) -> None:
        pages = list(self.doc_pdf)
        dicts, flags = [], []
        for i, p in enumerate(pages):
            d, o = self.page_dict(i, p)
            dicts.append(d); flags.append(o)
        if len(self.ocr_pages) < len(self.scanned) and self.ocr_lang and len(self.ocr_pages) >= self.ocr_max:
            self.warnings.append(f"OCR limited to the first {self.ocr_max} scanned pages; the rest were inserted as images.")
        self.body = body_font_size(dicts)
        for i, page in enumerate(pages):
            self.add_page(i, page, dicts[i], flags[i])
        if self.ocr_pages:
            self.warnings.append("OCR text is a best-effort recognition; proofread it, especially numbers and names.")
        self.out.save(out_path)
        self.doc_pdf.close()

    # ---- section / margins -----------------------------------------------
    def new_section(self, i: int, page, items: list[dict]):
        if i == 0:
            sec = self.out.sections[0]
        else:
            self.out.add_section(WD_SECTION.NEW_PAGE)
            shrink_paragraph(self.out.paragraphs[-1])
            sec = self.out.sections[-1]
        w, h = min(page.rect.width, 1584), min(page.rect.height, 1584)
        sec.page_width, sec.page_height = Pt(w), Pt(h)
        sec.orientation = WD_ORIENT.LANDSCAPE if w > h else WD_ORIENT.PORTRAIT
        if items:
            l = min(i_["x0"] for i_ in items); r = max(i_["x1"] for i_ in items)
            t = min(i_["y0"] for i_ in items); b = max(i_["y1"] for i_ in items)
        else:
            l, r, t, b = 54, w - 54, 54, h - 54
        clamp = lambda v: max(28.0, min(v, 120.0))
        sec.left_margin, sec.right_margin = Pt(clamp(l)), Pt(clamp(w - r))
        sec.top_margin, sec.bottom_margin = Pt(clamp(t)), Pt(clamp(h - b))
        sec.header_distance = sec.footer_distance = Pt(0)
        return sec, clamp(l), clamp(w - r)

    # ---- one page --------------------------------------------------------
    def add_page(self, i: int, page, d: dict, ocr: bool) -> None:
        self.stats["pages"] += 1
        text_blocks = [b for b in d["blocks"] if b["type"] == 0 and any(s["text"].strip() for ln in b["lines"] for s in ln["spans"])]
        scanned_image_only = i in self.scanned and not ocr
        items: list[dict] = []

        tables = []
        if not scanned_image_only:
            try:
                tables = [t for t in page.find_tables().tables if t.row_count >= 2 and t.col_count >= 2]
            except Exception:
                tables = []
        tbboxes = [pymupdf.Rect(t.bbox) for t in tables]

        for b in text_blocks:
            r = pymupdf.Rect(b["bbox"])
            c = pymupdf.Point((r.x0 + r.x1) / 2, (r.y0 + r.y1) / 2)
            if any(c in tb for tb in tbboxes):
                continue
            items.append({"kind": "text", "x0": r.x0, "y0": r.y0, "x1": r.x1, "y1": r.y1, "lines": len(b["lines"]), "block": b})
        for t in tables:
            x0, y0, x1, y1 = t.bbox
            items.append({"kind": "table", "x0": x0, "y0": y0, "x1": x1, "y1": y1, "lines": t.row_count, "table": t})

        if scanned_image_only:
            items.append({"kind": "page_image", "x0": 0, "y0": 0, "x1": page.rect.width, "y1": page.rect.height, "lines": 0})
        elif not ocr:
            seen = set()
            page_area = page.rect.width * page.rect.height
            for im in page.get_image_info(xrefs=True):
                r = pymupdf.Rect(im["bbox"]) & page.rect
                key = (im.get("xref"), tuple(round(v) for v in r))
                if r.is_empty or r.width < 8 or r.height < 8 or key in seen or not im.get("xref"):
                    continue
                seen.add(key)
                if r.width * r.height > 0.85 * page_area and text_blocks:
                    self.warnings.append(f"Page {i + 1}: a full-page background image was left out so the text stays editable (use Maximum Layout Preservation to keep it).")
                    continue
                items.append({"kind": "image", "x0": r.x0, "y0": r.y0, "x1": r.x1, "y1": r.y1, "lines": 0, "xref": im["xref"]})
            try:
                nd = len(page.get_drawings())
            except Exception:
                nd = 0
            if nd > 60 and not tables:
                self.warnings.append(f"Page {i + 1}: contains vector graphics/shapes (charts, lines or diagrams) that Editable mode cannot reproduce. Use Maximum Layout Preservation for this page.")

        _, ml, mr = self.new_section(i, page, items)
        avail_w = page.rect.width - ml - mr
        prev_bottom = None
        for it in column_order(items, page.rect.width):
            gap = 0 if prev_bottom is None else max(0.0, it["y0"] - prev_bottom)
            if it["kind"] == "text":
                self.add_text_block(it, ml, mr, page.rect.width, min(gap, 40))
            elif it["kind"] == "table":
                self.add_table(page, it["table"], avail_w)
            elif it["kind"] == "image":
                self.add_image(it, avail_w)
            else:
                self.add_page_image(page, avail_w)
            prev_bottom = it["y1"]
        if not items:
            self.out.add_paragraph("")

    # ---- elements --------------------------------------------------------
    def add_page_image(self, page, avail_w: float) -> None:
        pix = page.get_pixmap(dpi=150)
        p = self.out.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.add_run().add_picture(io.BytesIO(pix.tobytes("png")), width=Pt(avail_w))
        self.stats["image_only_pages"] += 1
        self.stats["images"] += 1
        self.warnings.append(f"Page {page.number + 1} is a scanned image and was inserted as a picture (OCR was off or unavailable), so its text is not editable.")

    def add_image(self, it: dict, avail_w: float) -> None:
        data = extract_image_bytes(self.doc_pdf, it["xref"])
        if not data:
            self.warnings.append("An image could not be extracted and was skipped.")
            return
        w = min(it["x1"] - it["x0"], avail_w)
        p = self.out.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = pf.space_after = Pt(2)
        p.add_run().add_picture(io.BytesIO(data), width=Pt(max(w, 8)))
        self.stats["images"] += 1

    @staticmethod
    def table_rows(page, t) -> list[list]:
        """Read each cell's text straight from its rectangle (keeps words/decimals intact)."""
        rows = []
        for r in t.rows:
            row = []
            for c in r.cells:
                if c is None:
                    row.append(None)
                else:
                    rect = pymupdf.Rect(c) + (1, 1, -1, -1)
                    row.append(page.get_text("text", clip=rect).strip())
            rows.append(row)
        return rows or t.extract()

    def add_table(self, page, t, avail_w: float) -> None:
        rows = self.table_rows(page, t)
        ncols = max(len(r) for r in rows)
        table = self.out.add_table(rows=len(rows), cols=ncols)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        widths = None
        try:
            first = t.rows[0].cells
            if len(first) == ncols and all(c is not None for c in first):
                widths = [c[2] - c[0] for c in first]
        except Exception:
            pass
        if not widths:
            widths = [(t.bbox[2] - t.bbox[0]) / ncols] * ncols
        scale = min(1.0, avail_w / (sum(widths) or 1))
        for ri, row in enumerate(rows):
            try:  # keep the PDF's row height (as a minimum) so tables don't look cramped
                rh = t.rows[ri].bbox[3] - t.rows[ri].bbox[1]
                table.rows[ri].height = Pt(max(rh, 12))
                table.rows[ri].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
            except Exception:
                pass
            cells = table.rows[ri].cells
            last_real = None
            for ci in range(ncols):
                val = row[ci] if ci < len(row) else None
                cell = cells[ci]
                cell.width = Pt(widths[ci] * scale)
                if val is None and last_real is not None:
                    last_real.merge(cell)  # None = continuation of a merged cell
                    continue
                last_real = cell
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                para = cell.paragraphs[0]
                para.paragraph_format.space_after = Pt(0)
                text = clean((val or "").strip())
                lines = text.split("\n")
                for li, line in enumerate(lines):
                    run = para.add_run(line)
                    run.font.size = Pt(10)
                    set_run_fonts(run, "Calibri", bool(BENGALI_RE.search(line)))
                    if li < len(lines) - 1:
                        run.add_break()
        self.out.add_paragraph().paragraph_format.space_after = Pt(0)
        self.stats["tables"] += 1

    def add_text_block(self, it: dict, ml: float, mr: float, page_w: float, gap: float) -> None:
        b = it["block"]
        lines = [ln for ln in b["lines"] if any(s["text"].strip() for s in ln["spans"])]
        if not lines:
            return
        bw = max(b["bbox"][2] - b["bbox"][0], 1)
        sizes = [s["size"] for ln in lines for s in ln["spans"] if s["text"].strip()]
        top_size = max(sizes)
        all_bold = all((s["flags"] & 16) or "bold" in s["font"].lower() for ln in lines for s in ln["spans"] if s["text"].strip())
        first_text = "".join(s["text"] for s in lines[0]["spans"])
        heading = None
        if len(lines) <= 3 and top_size >= self.body * 1.25 and (all_bold or top_size >= self.body * 1.5):
            heading = 1 if top_size >= self.body * 1.8 else 2 if top_size >= self.body * 1.4 else 3
        para = self.out.add_heading(level=heading) if heading else self.out.add_paragraph()
        if heading:
            self.stats["headings"] += 1
        self.stats["paragraphs"] += 1
        pf = para.paragraph_format
        pf.space_before, pf.space_after = Pt(gap), Pt(0)
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        pf.line_spacing = 1.0
        # alignment & indent
        x0, x1 = b["bbox"][0], b["bbox"][2]
        left_gap, right_gap = x0 - ml, (page_w - mr) - x1
        if abs(left_gap - right_gap) < 10 and left_gap > 24:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif right_gap < 6 and left_gap > 80 and len(lines) <= 2:
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        else:
            if left_gap > 12:
                pf.left_indent = Pt(min(left_gap, 300))
        if BULLET_RE.match(first_text) or NUMBER_RE.match(first_text):
            pf.left_indent = Pt(max(min(left_gap, 300), 0) + 18)
            pf.first_line_indent = Pt(-18)
        # runs
        for li, ln in enumerate(lines):
            spans = [s for s in ln["spans"] if s["text"] != ""]
            for s in spans:
                text = clean(s["text"])
                if not text:
                    continue
                run = para.add_run(text)
                run.font.size = Pt(round(s["size"] * 2) / 2 or 11)
                bold = bool(s["flags"] & 16) or "bold" in s["font"].lower()
                italic = bool(s["flags"] & 2) or any(k in s["font"].lower() for k in ("italic", "oblique"))
                run.bold, run.italic = bold, italic
                if s["flags"] & 1:
                    run.font.superscript = True
                col = s.get("color", 0)
                run.font.color.rgb = RGBColor((col >> 16) & 255, (col >> 8) & 255, col & 255)
                set_run_fonts(run, map_font(s["font"]), bool(BENGALI_RE.search(text)))
            if li < len(lines) - 1:
                lw = ln["bbox"][2] - ln["bbox"][0]
                if lw < 0.82 * bw or ln["bbox"][3] - ln["bbox"][1] < 0:
                    para.add_run().add_break()  # short line = intentional line break
                else:
                    para.add_run(" ")


def build_editable(pdf_path: str, out_path: str, *, ocr_lang: str | None, ocr_tessdata: str | None,
                   ocr_max_pages: int, scanned_pages: list[int]) -> dict:
    b = Builder(pdf_path, ocr_lang=ocr_lang, ocr_tessdata=ocr_tessdata, ocr_max_pages=ocr_max_pages,
                scanned_pages={p - 1 for p in scanned_pages})
    b.build(out_path)
    return {"warnings": b.warnings, "stats": b.stats, "ocr_pages": b.ocr_pages}
