import pytest
from docx import Document

from app.analyzer import analyze_pdf
from app.converter import run_job
from app.ocr import ocr_status, resolve_languages
from app.quality import docx_stats, render_page_count

OCR = ocr_status()
needs_ocr = pytest.mark.skipif("eng" not in OCR["languages"], reason="Tesseract English data not installed")


def convert(pdf, tmp_path, mode, ocr=True, lang="eng"):
    out = tmp_path / f"{mode}.docx"
    report = run_job(str(pdf), str(out), mode, ocr, lang, 30)
    return out, report


def check(report, name):
    return next(c for c in report["checks"] if c["name"] == name)


def test_analyzer_detects_scanned(text_pdf, scanned_pdf, table_image_pdf):
    assert analyze_pdf(str(text_pdf))["scanned_pages"] == []
    assert analyze_pdf(str(scanned_pdf))["scanned_pages"] == [1]
    a = analyze_pdf(str(table_image_pdf))
    assert a["tables"] >= 1 and a["images"] >= 1


def test_editable_text_pages_and_headings(text_pdf, tmp_path):
    out, rep = convert(text_pdf, tmp_path, "editable")
    doc = Document(out)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Quarterly Report" in text and "12345" in text and "Second bullet point" in text
    assert any(p.style.name.startswith("Heading") and "Quarterly" in p.text for p in doc.paragraphs)
    assert len(doc.sections) == 2  # one section per PDF page
    assert check(rep, "text_coverage")["value"] >= 97


def test_editable_table_and_image(table_image_pdf, tmp_path):
    out, rep = convert(table_image_pdf, tmp_path, "editable")
    doc = Document(out)
    assert len(doc.tables) == 1
    cells = [[c.text for c in r.cells] for r in doc.tables[0].rows]
    assert cells[0] == ["Item", "Qty", "Price"] and cells[2] == ["Banana", "24", "0.75"]
    assert docx_stats(str(out))["images"] >= 1
    assert check(rep, "tables")["status"] == "pass" and check(rep, "images")["status"] == "pass"


def test_preserve_mode(table_image_pdf, tmp_path):
    out, rep = convert(table_image_pdf, tmp_path, "preserve")
    st = docx_stats(str(out))
    assert st["images"] >= 1 and st["tables"] >= 1
    assert check(rep, "text_coverage")["value"] >= 90
    assert rep["mode"] == "preserve"


def test_preserve_text_document(text_pdf, tmp_path):
    out, rep = convert(text_pdf, tmp_path, "preserve")
    assert check(rep, "text_coverage")["value"] >= 90


@needs_ocr
def test_scanned_editable_ocr(scanned_pdf, tmp_path):
    out, rep = convert(scanned_pdf, tmp_path, "editable", ocr=True, lang="eng")
    text = " ".join(p.text for p in Document(out).paragraphs).upper()
    assert "INVOICE" in text and "BUSINESS" in text
    assert rep["ocr"]["used"] is True
    assert any("proofread" in w for w in rep["warnings"])


def test_scanned_without_ocr_becomes_image(scanned_pdf, tmp_path):
    out, rep = convert(scanned_pdf, tmp_path, "editable", ocr=False)
    assert docx_stats(str(out))["images"] == 1
    assert rep["level"] == "review" and any("image" in w.lower() for w in rep["warnings"])
    assert check(rep, "text_coverage")["status"] == "skip"


def test_scanned_preserve_is_honest(scanned_pdf, tmp_path):
    out, rep = convert(scanned_pdf, tmp_path, "preserve")
    assert docx_stats(str(out))["images"] >= 1
    assert any("not editable" in w for w in rep["warnings"])


def test_missing_bangla_data_is_reported_not_faked():
    langs, warns = resolve_languages("eng+ben", {"languages": ["eng"]})
    assert langs == "eng" and warns and "Bangla" in warns[0]
    langs, warns = resolve_languages("ben", {"languages": ["eng"]})
    assert langs is None and "skipped" in warns[0]


def test_bengali_run_font_is_set():
    from docx import Document as D
    from app.editable import set_run_fonts
    run = D().add_paragraph().add_run("বাংলা")
    set_run_fonts(run, "Arial", True)
    assert 'w:cs="Nirmala UI"' in run._element.xml


@pytest.mark.skipif(render_page_count.__module__ and not __import__("shutil").which("soffice"), reason="LibreOffice not installed")
def test_libreoffice_page_count(text_pdf, tmp_path, monkeypatch):
    monkeypatch.setenv("RUDRIX_QUALITY_RENDER", "true")
    out, _ = convert(text_pdf, tmp_path, "editable")
    assert render_page_count(str(out)) == 2
